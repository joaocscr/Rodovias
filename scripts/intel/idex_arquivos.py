"""Leitor dos arquivos do IDEX-Infra (data/raw/idex/, fora do git).

É a única parte do domínio que conhece o formato dos arquivos. Para ler da API do IDEX no futuro, escreva um módulo
irmão com as mesmas funções e o mesmo contrato de saída e passe-as a `idex.montar`.

- `arquivos(pasta)` → {"datafiles": [Path], "relatorios": [Path]}. `datafiles`: todos os `*datafile*.xlsx` da pasta, por
  nome (a data de modificação não conta: cópias e sincronização a mudam); `relatorios`: os `*.pdf` da pasta, por nome.
- `ler_datafile(path)` → [{data: "AAAA-MM-DD", ticker, peso: float, duration: float | None, spread_bps: float | None}].
  Lê a aba "Detalhado", colunas `Data`, `Debênture`, `Peso no índice (%)`, `Duration` e `MID spread (Bps/NTNB)`. O peso é
  uma fração (a soma do dia é 1), apesar do "(%)" do cabeçalho. Aba ou coluna ausente: ValueError.
- `ler_relatorio(path)` → {"mes_referencia": "AAAA-MM", "retornos": {nome: {"AAAA-MM": pct}},
  "acumulado_ano": {nome: {"AAAA": pct}}}. Lê a página "Histórico de performance" (pymupdf): para cada ano, uma linha
  por índice ("Idex-Infra Geral", "Idex-Infra Core") com os retornos mensais a partir de janeiro e, por último, o
  "Acum. Ano" (um ano incompleto tem menos de doze meses). `mes_referencia` é o último mês com retorno. Página ausente
  ou tabela que não se reconhece: ValueError.
"""
import datetime as dt
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

ABA = "Detalhado"
COLUNAS = {"data": "Data", "ticker": "Debênture", "peso": "Peso no índice (%)", "duration": "Duration",
           "spread_bps": "MID spread (Bps/NTNB)"}
PAGINA = "Histórico de performance"
# um número ("3,30 %", "-1,85%", "0,27 % "), um nome de índice ou um ano isolado, na ordem em que aparecem no texto
TOKEN = re.compile(r"(?P<pct>[-−–]?\d+,\d+\s*%)|(?P<nome>Idex-Infra\s+(?:Geral|Core))|(?<![\d,.])(?P<ano>20\d\d)(?![\d,%])",
                   re.I)


def arquivos(pasta):
    pasta = Path(pasta)
    datafiles = sorted(p for p in pasta.glob("*datafile*.xlsx") if not p.name.startswith("~$"))
    return {"datafiles": datafiles, "relatorios": sorted(pasta.glob("*.pdf"))}


def _numero(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _dia(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v or "").strip()
    try:
        return dt.date.fromisoformat(s[:10]).isoformat()
    except ValueError:
        return None


def ler_datafile(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if ABA not in wb.sheetnames:
            raise ValueError(f"{path}: sem a aba \"{ABA}\"")
        linhas = wb[ABA].iter_rows(values_only=True)
        cab = [str(c).strip() if c is not None else None for c in next(linhas, ())]
        falta = [c for c in COLUNAS.values() if c not in cab]
        if falta:
            raise ValueError(f"{path}: colunas ausentes na aba \"{ABA}\": {', '.join(falta)}")
        pos = {k: cab.index(c) for k, c in COLUNAS.items()}
        out = []
        for r in linhas:
            if len(r) <= max(pos.values()):
                continue
            data, ticker, peso = _dia(r[pos["data"]]), str(r[pos["ticker"]] or "").strip().upper(), _numero(r[pos["peso"]])
            if not data or not ticker or peso is None:
                continue
            out.append({"data": data, "ticker": ticker, "peso": peso, "duration": _numero(r[pos["duration"]]),
                        "spread_bps": _numero(r[pos["spread_bps"]])})
        return out
    finally:
        wb.close()


def _pct(txt):
    return float(re.sub(r"[−–]", "-", txt).replace("%", "").replace(",", ".").strip())


def _linhas_da_tabela(texto):
    """[(ano, nome, [valores])] do texto da página, na ordem em que aparecem."""
    ano = nome = None
    linhas = []
    for m in TOKEN.finditer(texto):
        if m["ano"]:
            ano, nome = int(m["ano"]), None
        elif m["nome"]:
            if ano is None:
                raise ValueError("índice antes de qualquer ano na tabela")
            nome = "Idex-Infra " + m["nome"].split()[-1].capitalize()
            linhas.append((ano, nome, []))
        else:
            if nome is None:
                raise ValueError("número fora de uma linha de índice na tabela")
            linhas[-1][2].append(_pct(m["pct"]))
    return linhas


def ler_relatorio(path):
    import pymupdf
    with pymupdf.open(path) as doc:
        textos = [p.get_text() for p in doc]
    pagina = next((t for t in textos if PAGINA.lower() in t.lower()), None)
    if pagina is None:
        raise ValueError(f"{path}: sem a página \"{PAGINA}\"")
    texto = pagina[pagina.lower().index(PAGINA.lower()) + len(PAGINA):]
    linhas = _linhas_da_tabela(texto)
    if not linhas:
        raise ValueError(f"{path}: a tabela de \"{PAGINA}\" não foi reconhecida")
    retornos, acumulado, ultimo = {}, {}, None
    for ano, nome, valores in linhas:
        if len(valores) < 2 or len(valores) > 13:
            raise ValueError(f"{path}: {nome} em {ano} traz {len(valores)} números (esperado de 2 a 13: meses e acumulado)")
        if str(ano) in acumulado.get(nome, {}):
            raise ValueError(f"{path}: {nome} aparece duas vezes em {ano}")
        meses, acum = valores[:-1], valores[-1]  # alinhados de janeiro em diante; o último número é o "Acum. Ano"
        for i, v in enumerate(meses, start=1):
            retornos.setdefault(nome, {})[f"{ano}-{i:02d}"] = v
        acumulado.setdefault(nome, {})[str(ano)] = acum
        ultimo = max(ultimo or "", f"{ano}-{len(meses):02d}")
    return {"mes_referencia": ultimo, "retornos": retornos, "acumulado_ano": acumulado}
